"""Builds a .pptx from a deck folder: deck.yaml (the slides) + style.md (the look).

Every slide is drawn from the style's tokens on PowerPoint's "Title Only" or
"Blank" layout, so the same deck.yaml always gives the same deck. The theme's
colours and fonts are set too, so a slide added by hand in PowerPoint matches.
"""
from __future__ import annotations

import hashlib
import logging
import math
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .style import load_style, load_yaml

log = logging.getLogger("presentations")

W, H = 13.333, 7.5  # 16:9 in inches
M = 0.6  # outer margin
TITLE_Y, TITLE_H = 0.55, 1.15
TOP, BOTTOM = 1.85, 6.6  # content area under the title
GAP = 0.35

LAYOUT_TITLE_ONLY, LAYOUT_BLANK = 5, 6
PICTURE_TYPES = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff"}
ANCHORS = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
ALIGNS = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}

Box = Tuple[float, float, float, float]


class Ctx:
    """What a layout needs besides its own slide data."""

    def __init__(self, style: Dict[str, Any], deck_dir: Path):
        self.st = style
        self.c = style["colors"]
        self.f = style["fonts"]
        self.s = style["sizes"]
        self.deck_dir = deck_dir
        self.label = ""  # "slide 3 (bullets)", for warnings


# ---------- small drawing helpers ----------

def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def mix(a: str, b: str, t: float) -> str:
    """Colour t of the way from a to b, for tints the style doesn't name."""
    ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))


def background(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def rect(slide, box: Box, color: str, shape=MSO_SHAPE.RECTANGLE, name: str = ""):
    x, y, w, h = box
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def add_runs(p, text: str, font: str, size: float, color: str, bold=False, italic=False) -> None:
    # **word** turns bold, the only inline markup deck.yaml supports.
    for i, part in enumerate(re.split(r"\*\*", str(text))):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold or i % 2 == 1
        r.font.italic = italic
        r.font.color.rgb = rgb(color)


def frame(tf, anchor: str = "top") -> None:
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHORS[anchor]


def estimate_height(lines: List[Tuple[str, float]], width: float, ratio: float = 0.52) -> float:
    """Rough height in inches of wrapped text: ratio is average glyph width / font size."""
    total = 0.0
    for text, size in lines:
        per_line = max(1, int(width * 72 / (size * ratio)))
        total += max(1, math.ceil(len(text) / per_line)) * size * 1.2 / 72
    return total


def warn_fit(ctx: Ctx, what: str, lines: List[Tuple[str, float]], box: Box, extra: float = 0.0) -> None:
    need = estimate_height(lines, box[2]) + extra
    if need > box[3] * 1.05:
        log.warning("%s: %s probably overflows (needs ~%.1fin, has %.1fin). Shorten it or split the slide.",
                    ctx.label, what, need, box[3])


def text(slide, box: Box, value: str, ctx: Ctx, *, size: float, color: str, font: Optional[str] = None,
         bold=False, italic=False, align="left", anchor="top", name: str = "", check: bool = True) -> None:
    x, y, w, h = box
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    frame(tf, anchor)
    for i, para in enumerate(str(value).strip().split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGNS[align]
        add_runs(p, para, font or ctx.f["body"], size, color, bold, italic)
    if check:
        warn_fit(ctx, name or "text", [(line, size) for line in str(value).split("\n")], box)


def flatten_bullets(items: List[Any]) -> List[Tuple[str, int]]:
    """deck.yaml bullets: a string, or a one-key mapping {point: [sub-points]}."""
    out: List[Tuple[str, int]] = []
    for item in items or []:
        if isinstance(item, dict):
            for head, subs in item.items():
                # An unquoted "Key: value" bullet parses as a mapping to a scalar,
                # not to sub-points; put the text back together.
                if subs is not None and not isinstance(subs, list):
                    out.append((f"{head}: {subs}", 0))
                    continue
                out.append((str(head), 0))
                out.extend((str(s), 1) for s in subs or [])
        else:
            out.append((str(item), 0))
    return out


def bullets(slide, box: Box, items: List[Any], ctx: Ctx, *, size: Optional[float] = None,
            color: Optional[str] = None, intro: str = "") -> None:
    size = size or ctx.s["body"]
    color = color or ctx.c["text"]
    x, y, w, h = box
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = "Bullets"
    tf = tb.text_frame
    frame(tf)
    first = True
    if intro:
        p = tf.paragraphs[0]
        p.space_after = Pt(size * 0.8)
        add_runs(p, intro, ctx.f["body"], size, color)
        first = False
    rows = flatten_bullets(items)
    for line, level in rows:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        sub = size - 2 if level else size
        p.space_after = Pt(sub * 0.6)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.32 + 0.35 * level)))
        pPr.set("indent", str(-Inches(0.32)))
        # Real bullets (buChar), not a typed "•", so PowerPoint treats it as a list.
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr"), val=ctx.c["accent"])
        etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
        etree.SubElement(pPr, qn("a:buChar"), char="–" if level else "•")
        add_runs(p, line, ctx.f["body"], sub, color)
    lines = ([(intro, size)] if intro else []) + [(l, size - 2 if lv else size) for l, lv in rows]
    warn_fit(ctx, "bullets", lines, (x + 0.32, y, w - 0.32, h), extra=len(rows) * size * 0.6 / 72)


def resolve_image(ctx: Ctx, value: str) -> Optional[Path]:
    path = Path(str(value)).expanduser()
    if not path.is_absolute():
        path = ctx.deck_dir / path
    if not path.exists():
        return None
    if path.suffix.lower() in PICTURE_TYPES:
        return path
    # WebP etc.: python-pptx only embeds the formats above, so convert once.
    out = ctx.deck_dir / ".converted" / (path.stem + ".png")
    if not out.exists() or out.stat().st_mtime < path.stat().st_mtime:
        out.parent.mkdir(exist_ok=True)
        with Image.open(path) as im:
            im.convert("RGB").save(out)
    return out


def picture(slide, box: Box, value: Optional[str], ctx: Ctx) -> None:
    """Fill box with the image, cropped (not squashed) to the box's shape."""
    x, y, w, h = box
    path = resolve_image(ctx, value) if value else None
    if path is None:
        log.warning("%s: image %r not found; drew a placeholder.", ctx.label, value)
        rect(slide, box, ctx.c["surface"], name="Image placeholder")
        text(slide, (x + 0.4, y + h / 2 - 0.3, w - 0.8, 0.6), f"Billede mangler: {value or '(intet)'}",
             ctx, size=ctx.s["caption"], color=ctx.c["muted"], align="center", anchor="middle")
        return
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.name = "Image"
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:
        pic.crop_left = pic.crop_right = (1 - box_ratio / img_ratio) / 2
    else:
        pic.crop_top = pic.crop_bottom = (1 - img_ratio / box_ratio) / 2


def logo(slide, ctx: Ctx, box: Box) -> None:
    if not ctx.st.get("logo"):
        return
    path = ctx.st["dir"] / ctx.st["logo"]
    if not path.exists():
        log.warning("logo %s not found.", path)
        return
    x, y, w, h = box
    with Image.open(path) as im:
        ratio = im.width / im.height
    lw = min(w, h * ratio)
    slide.shapes.add_picture(str(path), Inches(x + w - lw), Inches(y), Inches(lw), Inches(lw / ratio)).name = "Logo"


def motif(slide, ctx: Ctx, on: str) -> None:
    """The style's one repeated decoration, on dark slides only."""
    if ctx.st.get("motif") == "circle":
        rect(slide, (W - 4.6, -1.6, 7.0, 7.0), mix(on, ctx.c["on_primary"], 0.07), MSO_SHAPE.OVAL, "Motif")
        rect(slide, (W - 2.15, 5.55, 0.55, 0.55), ctx.c["accent"], MSO_SHAPE.OVAL, "Motif dot")


def title(slide, value: str, ctx: Ctx, box: Optional[Box] = None, *, size: Optional[float] = None,
          color: Optional[str] = None, anchor: str = "top") -> None:
    """Fill the layout's title placeholder, so PowerPoint's outline and screen readers see it."""
    size = size or ctx.s["title"]
    x, y, w, h = box or (M, TITLE_Y, W - 2 * M, TITLE_H)
    ph = slide.shapes.title
    ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
    tf = ph.text_frame
    frame(tf, anchor)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    add_runs(p, value, ctx.f["heading"], size, color or ctx.c["text"])
    for extra in str(value).split("\n")[1:]:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        add_runs(p, extra, ctx.f["heading"], size, color or ctx.c["text"])
    warn_fit(ctx, "title", [(line, size) for line in str(value).split("\n")], (x, y, w, h))


def chrome(slide, ctx: Ctx, number: int, *, show_number: bool = True) -> None:
    """Footer text and slide number on content slides."""
    y = 6.95
    if ctx.st.get("footer"):
        text(slide, (M, y, 7.0, 0.3), ctx.st["footer"], ctx, size=ctx.s["caption"], color=ctx.c["muted"],
             name="Footer")
    if ctx.st.get("slide_numbers") and show_number:
        text(slide, (W - M - 1.0, y, 1.0, 0.3), str(number), ctx, size=ctx.s["caption"],
             color=ctx.c["muted"], align="right", name="Slide number")


def panel(slide, box: Box, ctx: Ctx):
    return rect(slide, box, ctx.c["surface"], MSO_SHAPE.RECTANGLE, "Panel")


# ---------- layouts ----------
# Each takes (slide, data, ctx, number). Keys a layout reads are listed in
# presentations/sample.yaml, which uses every layout once.

def lay_title(slide, d, ctx: Ctx, n: int) -> None:
    c = ctx.c
    background(slide, c["primary"])
    motif(slide, ctx, c["primary"])
    soft = mix(c["on_primary"], c["primary"], 0.3)
    # The title grows upwards from a fixed baseline, so the kicker follows its top.
    title_h = estimate_height([(t, ctx.s["cover"]) for t in str(d.get("title", "")).split("\n")], 8.8)
    if d.get("kicker"):
        text(slide, (M, 4.5 - title_h - 0.55, 8.5, 0.4), d["kicker"].upper(), ctx, size=ctx.s["caption"] + 2,
             color=c["accent"], bold=True, name="Kicker")
    title(slide, d.get("title", ""), ctx, (M, 2.2, 8.8, 2.3), size=ctx.s["cover"], color=c["on_primary"],
          anchor="bottom")
    if d.get("subtitle"):
        text(slide, (M, 4.75, 8.8, 1.1), d["subtitle"], ctx, size=ctx.s["heading"], color=soft, name="Subtitle")
    if d.get("byline"):
        text(slide, (M, 6.45, 8.0, 0.4), d["byline"], ctx, size=ctx.s["caption"] + 2, color=soft, name="Byline")
    logo(slide, ctx, (W - M - 2.0, 6.25, 2.0, 0.65))


def lay_section(slide, d, ctx: Ctx, n: int) -> None:
    c = ctx.c
    background(slide, c["primary"])
    motif(slide, ctx, c["primary"])
    size = ctx.s["cover"] - 4
    if d.get("number") is not None:
        text(slide, (M, 2.0, 3.0, 0.7), str(d["number"]), ctx, size=ctx.s["title"] - 6, color=c["accent"],
             font=ctx.f["heading"], name="Section number")
    title(slide, d.get("title", ""), ctx, (M, 2.75, 9.5, 1.9), size=size, color=c["on_primary"])
    if d.get("subtitle"):
        title_h = estimate_height([(t, size) for t in str(d.get("title", "")).split("\n")], 9.5)
        text(slide, (M, 2.75 + min(title_h, 1.9) + 0.3, 9.0, 1.0), d["subtitle"], ctx, size=ctx.s["heading"] - 2,
             color=mix(c["on_primary"], c["primary"], 0.3), name="Subtitle")


def lay_bullets(slide, d, ctx: Ctx, n: int) -> None:
    background(slide, ctx.c["background"])
    if d.get("image"):
        split = 7.6  # text left of this, image bleeds off the right edge
        title(slide, d.get("title", ""), ctx, (M, TITLE_Y, split - M - 0.5, TITLE_H))
        bullets(slide, (M, TOP, split - M - 0.5, BOTTOM - TOP), d.get("bullets"), ctx, intro=d.get("text", ""))
        picture(slide, (split, 0, W - split, H), d["image"], ctx)
        chrome(slide, ctx, n, show_number=False)
    else:
        title(slide, d.get("title", ""), ctx)
        bullets(slide, (M, TOP, 10.4, BOTTOM - TOP), d.get("bullets"), ctx, intro=d.get("text", ""))
        chrome(slide, ctx, n)


def column_height(d: Dict[str, Any], w: float, ctx: Ctx, pad: float) -> float:
    size = ctx.s["body"] - 1
    lines = ([(d["text"], size)] if d.get("text") else []) + [
        (t, size - 2 if lv else size) for t, lv in flatten_bullets(d.get("bullets"))]
    rows = len(lines)
    return (2 * pad + (0.75 if d.get("heading") else 0)
            + estimate_height(lines, w - 2 * pad - 0.32) + rows * size * 0.6 / 72)


def column(slide, box: Box, d: Dict[str, Any], ctx: Ctx) -> None:
    x, y, w, h = box
    pad = 0.4
    panel(slide, box, ctx)
    inner_y = y + pad
    if d.get("heading"):
        text(slide, (x + pad, inner_y, w - 2 * pad, 0.5), d["heading"], ctx, size=ctx.s["heading"],
             color=ctx.c["heading"], font=ctx.f["heading"], name="Column heading")
        inner_y += 0.75
    bullets(slide, (x + pad, inner_y, w - 2 * pad, y + h - pad - inner_y), d.get("bullets"), ctx,
            size=ctx.s["body"] - 1, intro=d.get("text", ""))


def fit_panel(need: float) -> float:
    """Panels are as tall as their content (with a floor), not always the full content area."""
    return min(BOTTOM - TOP, max(2.8, need + 0.3))


def lay_two_column(slide, d, ctx: Ctx, n: int) -> None:
    background(slide, ctx.c["background"])
    title(slide, d.get("title", ""), ctx)
    w = (W - 2 * M - GAP) / 2
    left, right = d.get("left") or {}, d.get("right") or {}
    h = fit_panel(max(column_height(left, w, ctx, 0.4), column_height(right, w, ctx, 0.4)))
    column(slide, (M, TOP, w, h), left, ctx)
    column(slide, (M + w + GAP, TOP, w, h), right, ctx)
    chrome(slide, ctx, n)


def lay_cards(slide, d, ctx: Ctx, n: int) -> None:
    background(slide, ctx.c["background"])
    title(slide, d.get("title", ""), ctx)
    cards = d.get("cards") or []
    if not 1 <= len(cards) <= 4:
        log.warning("%s: cards fits 1-4 cards, got %d.", ctx.label, len(cards))
    count = max(1, min(len(cards), 4))
    w = (W - 2 * M - GAP * (count - 1)) / count
    pad = 0.35
    inner = w - 2 * pad
    hs, ts = ctx.s["heading"] - 2, ctx.s["body"] - 2

    def mark_of(i, card):
        return card.get("number", f"{i + 1:02d}" if d.get("numbered", True) else None)

    head_h = max([estimate_height([(c.get("heading", ""), hs)], inner) for c in cards[:4]] or [0]) + 0.25
    need = max([2 * pad + (0.85 if mark_of(i, c) else 0) + head_h + estimate_height([(c.get("text", ""), ts)], inner)
                for i, c in enumerate(cards[:4])] or [0])
    h = fit_panel(need)
    for i, card in enumerate(cards[:4]):
        x = M + i * (w + GAP)
        panel(slide, (x, TOP, w, h), ctx)
        y = TOP + pad
        mark = mark_of(i, card)
        if mark:
            text(slide, (x + pad, y, inner, 0.6), str(mark), ctx, size=ctx.s["title"] - 6,
                 color=ctx.c["accent"], font=ctx.f["heading"], name="Card number")
            y += 0.85
        # Same heading height on every card, so the body text lines up across the row.
        text(slide, (x + pad, y, inner, head_h), card.get("heading", ""), ctx, size=hs,
             color=ctx.c["text"], font=ctx.f["heading"], name="Card heading")
        y += head_h
        text(slide, (x + pad, y, inner, TOP + h - pad - y), card.get("text", ""), ctx,
             size=ts, color=ctx.c["muted"], name="Card text")
    chrome(slide, ctx, n)


def lay_stats(slide, d, ctx: Ctx, n: int) -> None:
    background(slide, ctx.c["background"])
    title(slide, d.get("title", ""), ctx)
    stats = (d.get("stats") or [])[:4]
    count = max(1, len(stats))
    w = (W - 2 * M - GAP * (count - 1)) / count
    y = TOP + 0.6
    for i, st in enumerate(stats):
        x = M + i * (w + GAP)
        text(slide, (x, y, w, 1.3), str(st.get("value", "")), ctx, size=ctx.s["stat"], color=ctx.c["accent"],
             font=ctx.f["heading"], anchor="bottom", name="Stat value")
        text(slide, (x, y + 1.45, w - 0.2, 1.6), st.get("label", ""), ctx, size=ctx.s["body"],
             color=ctx.c["text"], name="Stat label")
    if d.get("text"):
        text(slide, (M, y + 3.25, 10.4, 1.0), d["text"], ctx, size=ctx.s["body"], color=ctx.c["muted"],
             name="Stats note")
    if d.get("source"):
        text(slide, (M, BOTTOM - 0.05, 10.0, 0.3), d["source"], ctx, size=ctx.s["caption"],
             color=ctx.c["muted"], italic=True, name="Source")
    chrome(slide, ctx, n)


def lay_image(slide, d, ctx: Ctx, n: int) -> None:
    background(slide, ctx.c["background"])
    if d.get("full") or not d.get("title"):
        picture(slide, (0, 0, W, H), d.get("image"), ctx)
        return
    split = 5.6
    title(slide, d["title"], ctx, (M, TITLE_Y, split - M - 0.5, 2.0))
    if d.get("text"):
        text(slide, (M, 2.85, split - M - 0.5, 3.6), d["text"], ctx, size=ctx.s["body"], color=ctx.c["text"],
             name="Text")
    if d.get("caption"):
        text(slide, (M, 6.45, split - M - 0.5, 0.45), d["caption"], ctx, size=ctx.s["caption"],
             color=ctx.c["muted"], italic=True, name="Caption")
    picture(slide, (split, 0, W - split, H), d.get("image"), ctx)


def lay_quote(slide, d, ctx: Ctx, n: int) -> None:
    c = ctx.c
    background(slide, c["surface"])
    text(slide, (M, 0.7, 2.0, 1.6), "“", ctx, size=140, color=c["accent"], font=ctx.f["heading"], name="Quote mark",
         check=False)  # one glyph; its line height overshoots the box but the mark itself fits
    text(slide, (M + 0.1, 2.35, 10.8, 3.2), d.get("quote", ""), ctx, size=ctx.s["title"] - 4, color=c["text"],
         font=ctx.f["heading"], italic=True, name="Quote")
    if d.get("by"):
        text(slide, (M + 0.1, 5.75, 10.0, 0.5), f"— {d['by']}", ctx, size=ctx.s["body"], color=c["muted"],
             name="Quote by")


def lay_table(slide, d, ctx: Ctx, n: int) -> None:
    c = ctx.c
    background(slide, c["background"])
    title(slide, d.get("title", ""), ctx)
    columns = [str(x) for x in d.get("columns") or []]
    rows = [[str(v) for v in r] for r in d.get("rows") or []]
    if not columns:
        log.warning("%s: table has no columns.", ctx.label)
        return
    row_h = 0.55
    if (len(rows) + 1) * row_h > BOTTOM - TOP:
        log.warning("%s: %d rows don't fit; split the table over two slides.", ctx.label, len(rows))
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(M), Inches(TOP), Inches(W - 2 * M),
                                   Inches(row_h * (len(rows) + 1)))
    shape.name = "Table"
    table = shape.table
    size = ctx.s["body"] - 2
    for r, values in enumerate([columns] + rows):
        for col in range(len(columns)):
            cell = table.cell(r, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(c["primary"] if r == 0 else (c["background"] if r % 2 else c["surface"]))
            cell.margin_left = cell.margin_right = Inches(0.15)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            add_runs(p, values[col] if col < len(values) else "", ctx.f["body"], size,
                     c["on_primary"] if r == 0 else c["text"], bold=r == 0)
    chrome(slide, ctx, n)


def lay_agenda(slide, d, ctx: Ctx, n: int) -> None:
    background(slide, ctx.c["background"])
    title(slide, d.get("title", "Agenda"), ctx)
    items = d.get("items") or []
    per_col = 5
    cols = 1 if len(items) <= per_col else 2
    w = (W - 2 * M - GAP) / 2 if cols == 2 else 10.4
    row = (BOTTOM - TOP) / per_col
    for i, item in enumerate(items[:per_col * 2]):
        x = M + (i // per_col) * (w + GAP)
        y = TOP + (i % per_col) * row
        text(slide, (x, y, 0.9, row), f"{i + 1:02d}", ctx, size=ctx.s["heading"] + 2, color=ctx.c["accent"],
             font=ctx.f["heading"], name="Agenda number")
        text(slide, (x + 1.0, y + 0.04, w - 1.0, row - 0.1), item, ctx, size=ctx.s["heading"] - 2,
             color=ctx.c["text"], name="Agenda item")
    if len(items) > per_col * 2:
        log.warning("%s: agenda fits %d items, got %d.", ctx.label, per_col * 2, len(items))
    chrome(slide, ctx, n)


def lay_closing(slide, d, ctx: Ctx, n: int) -> None:
    c = ctx.c
    background(slide, c["primary"])
    motif(slide, ctx, c["primary"])
    soft = mix(c["on_primary"], c["primary"], 0.3)
    title(slide, d.get("title", "Tak"), ctx, (M, 2.2, 8.8, 2.0), size=ctx.s["cover"], color=c["on_primary"],
          anchor="bottom")
    if d.get("subtitle"):
        text(slide, (M, 4.45, 8.8, 1.0), d["subtitle"], ctx, size=ctx.s["heading"], color=soft, name="Subtitle")
    contact = d.get("contact")
    if contact:
        lines = "\n".join(contact) if isinstance(contact, list) else str(contact)
        text(slide, (M, 5.55, 8.0, 1.3), lines, ctx, size=ctx.s["body"] - 2, color=soft, name="Contact")
    logo(slide, ctx, (W - M - 2.0, 6.25, 2.0, 0.65))


LAYOUTS: Dict[str, Tuple[Callable, int]] = {
    "title": (lay_title, LAYOUT_TITLE_ONLY),
    "agenda": (lay_agenda, LAYOUT_TITLE_ONLY),
    "section": (lay_section, LAYOUT_TITLE_ONLY),
    "bullets": (lay_bullets, LAYOUT_TITLE_ONLY),
    "two-column": (lay_two_column, LAYOUT_TITLE_ONLY),
    "cards": (lay_cards, LAYOUT_TITLE_ONLY),
    "stats": (lay_stats, LAYOUT_TITLE_ONLY),
    "image": (lay_image, LAYOUT_TITLE_ONLY),
    "quote": (lay_quote, LAYOUT_BLANK),
    "table": (lay_table, LAYOUT_TITLE_ONLY),
    "closing": (lay_closing, LAYOUT_TITLE_ONLY),
}


# ---------- theme, build, preview ----------

def apply_theme(prs, style: Dict[str, Any]) -> None:
    """Write the style's colours and fonts into the deck theme, for hand-made slides."""
    c = style["colors"]
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    scheme = {"dk1": c["text"], "lt1": c["background"], "dk2": c["primary"], "lt2": c["surface"],
              "accent1": c["accent"], "accent2": c["primary"], "accent3": c["muted"],
              "accent4": mix(c["accent"], c["background"], 0.4), "accent5": mix(c["primary"], c["background"], 0.4),
              "accent6": c["text"]}
    clr = root.find(f".//{qn('a:clrScheme')}")
    clr.set("name", style.get("name", "Style"))
    for key, value in scheme.items():
        slot = clr.find(qn(f"a:{key}"))
        for child in list(slot):
            slot.remove(child)
        etree.SubElement(slot, qn("a:srgbClr"), val=value)
    fonts = root.find(f".//{qn('a:fontScheme')}")
    fonts.find(f"{qn('a:majorFont')}/{qn('a:latin')}").set("typeface", style["fonts"]["heading"])
    fonts.find(f"{qn('a:minorFont')}/{qn('a:latin')}").set("typeface", style["fonts"]["body"])
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def register_notes_master(prs) -> None:
    """python-pptx adds the notes master part but not its <p:notesMasterIdLst>
    entry. PowerPoint doesn't mind; Keynote calls the file an invalid archive."""
    root = prs.part._element
    if root.find(qn("p:notesMasterIdLst")) is not None:
        return
    rid = next((r.rId for r in prs.part.rels.values() if r.reltype == RT.NOTES_MASTER), None)
    if rid is None:
        return
    lst = etree.Element(qn("p:notesMasterIdLst"))
    etree.SubElement(lst, qn("p:notesMasterId")).set(qn("r:id"), rid)
    root.find(qn("p:sldMasterIdLst")).addnext(lst)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build(deck_file: Path, out_dir: Path, force: bool = False) -> Path:
    """Write {slug}.pptx to out_dir. The .rendered stamp stays in the deck folder."""
    deck_dir = deck_file.parent
    style_file = deck_dir / "style.md"
    if not style_file.exists():
        raise FileNotFoundError(f"{style_file} is missing. Every deck folder needs its own style.md.")
    style = load_style(style_file)
    deck = load_yaml(deck_file)
    slug = deck.get("slug") or deck_dir.name
    out = out_dir / f"{slug}.pptx"

    # Re-rendering replaces the file, so refuse if someone edited it in PowerPoint since.
    # No stamp means the file isn't this deck's render (out_dir is shared, e.g. the Desktop).
    stamp = deck_dir / ".rendered"
    if out.exists() and not force:
        if not stamp.exists():
            raise FileExistsError(f"{out} exists but wasn't rendered from {deck_dir.name}. "
                                  "Move or rename it, change slug in deck.yaml, or pass --force to overwrite it.")
        if digest(out) != stamp.read_text().strip():
            raise FileExistsError(f"{out} was changed after the last render (edited in PowerPoint?). "
                                  "Move those edits into deck.yaml, or pass --force to overwrite them.")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    apply_theme(prs, style)
    props = prs.core_properties
    props.title = deck.get("title", slug)
    props.author = deck.get("author", "")

    ctx = Ctx(style, deck_dir)
    for n, data in enumerate(deck.get("slides") or [], start=1):
        name = data.get("layout", "bullets")
        if name not in LAYOUTS:
            raise ValueError(f"slide {n}: unknown layout '{name}'. Known: {sorted(LAYOUTS)}")
        fn, base = LAYOUTS[name]
        ctx.label = f"slide {n} ({name})"
        slide = prs.slides.add_slide(prs.slide_layouts[base])
        fn(slide, data, ctx, n)
        if data.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(data["notes"]).strip()

    register_notes_master(prs)
    prs.save(out)
    stamp.write_text(digest(out))
    log.info("wrote %s (%d slides)", out, len(prs.slides))
    return out


PREVIEW_SCRIPT = """
on run argv
  set src to POSIX file (item 1 of argv)
  set dest to POSIX file (item 2 of argv)
  set wasRunning to application "Keynote" is running
  tell application "Keynote"
    launch
    delay 2
    set doc to open src
    export doc to dest as slide images with properties {image format:PNG, skipped slides:false}
    close doc saving no
    if not wasRunning then quit
  end tell
end run
"""


def preview(pptx: Path, out: Path) -> List[Path]:
    """Slide PNGs in out (the deck's preview/), via Keynote (no LibreOffice on this Mac).

    Keynote's rendering is close to PowerPoint's but not identical; fonts that
    aren't installed here are substituted in the preview only.
    """
    shutil.rmtree(out, ignore_errors=True)
    result = subprocess.run(["osascript", "-e", PREVIEW_SCRIPT, str(pptx), str(out)],
                            capture_output=True, text=True, timeout=300)
    if result.returncode != 0:
        log.warning("preview failed (Keynote): %s", result.stderr.strip())
        return []
    pngs = sorted(out.glob("*.png"))
    renamed = []
    for i, png in enumerate(pngs, start=1):
        target = out / f"slide-{i:02d}.png"
        png.rename(target)
        renamed.append(target)
    return renamed
